#!/usr/bin/env python3
"""Actualiza el TDR con la arquitectura nativa de ATLAS de septiembre de 2026."""

from __future__ import annotations

import argparse
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent
DOCUMENT = ROOT / "atlas-teorico-codex.docx"
BACKUP = ROOT / "atlas-teorico-codex.pre-native-20260914.docx"


REPLACEMENTS: dict[int, tuple[str, str]] = {
    30: ("OpenClaw como base", "Primer prototipo con OpenClaw"),
    31: ("Funcionamiento de OpenClaw", "Qué aportó OpenClaw y por qué se descartó"),
    32: ("Adaptación de OpenClaw", "Migración del contexto a ATLAS nativo"),
    36: ("Gateway, sesiones", "Native Broker, sesiones y canales de comunicación"),
    37: ("Contexto, memoria", "Contexto, memoria y archivos de .atlas"),
    111: (
        "Para esta base decidí utilizar OpenClaw",
        "Para construir el primer prototipo decidí utilizar OpenClaw, una herramienta creada por "
        "el ingeniero y desarrollador de software austriaco Peter Steinberger.",
    ),
    112: (
        "OpenClaw permite crear agentes",
        "OpenClaw permitía conectar modelos, memoria, herramientas y canales con bastante rapidez, "
        "por lo que fue una buena forma de demostrar la idea inicial. Más adelante lo retiré del "
        "producto porque su cadena de procesos añadía demasiada latencia para una conversación por voz.",
    ),
    132: (
        "Y dado que el proyecto estaba",
        "Durante esa primera etapa, como el proyecto estaba potenciado en gran parte por OpenClaw, "
        "consideré llamar “OpenAtlas” al conjunto. El nombre terminó quedándose simplemente en ATLAS "
        "cuando el sistema empezó a tener una arquitectura y una identidad propias.",
    ),
    137: (
        "Una vez preparado ya el requisito",
        "Una vez preparado el primer prototipo con OpenClaw, ya podía comenzar a pensar cómo convertir "
        "la idea en un dispositivo real. De ahí surgió este planteamiento general:",
    ),
    139: (
        "OpenClaw será la base",
        "La primera versión utilizaría OpenClaw para unir modelos, contexto y herramientas. El diseño "
        "actual sustituye esa base por servicios propios de ATLAS, Realtime directo y un broker local.",
    ),
    150: ("2. OpenClaw como base", "2. Primer prototipo con OpenClaw y migración a ATLAS nativo"),
    152: (
        "Para desarrollar ATLAS era necesario",
        "Para desarrollar ATLAS era necesario conectar modelos de lenguaje, memoria, herramientas y "
        "canales. OpenClaw fue la primera base que utilicé para validar ese conjunto antes de construir "
        "un runtime propio más rápido y ajustado al dispositivo.",
    ),
    154: (
        "En este apartado voy a explicar",
        "En este apartado explico qué aportó OpenClaw al prototipo, qué problemas aparecieron y qué "
        "partes conservé al migrar ATLAS a su arquitectura actual.",
    ),
    157: ("2.1 Funcionamiento", "2.1 Funcionamiento del primer prototipo con OpenClaw"),
    159: (
        "Para poder entender cómo funciona ATLAS",
        "Para entender la evolución de ATLAS conviene explicar primero cómo funcionaba aquel prototipo "
        "y qué responsabilidades resolvía OpenClaw.",
    ),
    171: (
        "Pero para poder entender realmente",
        "Para entender aquella etapa no basta con mirar la página de conversación. También hay que "
        "revisar cómo OpenClaw distribuía la configuración, la memoria y las instrucciones del agente.",
    ),
    173: (
        "Tras finalizar la configuración",
        "Al terminar la configuración se creaba la carpeta .openclaw dentro del directorio del usuario. "
        "En el prototipo, allí convivían la configuración, las credenciales, las sesiones y el workspace "
        "del agente. Esta estructura resultaba cómoda para comenzar, pero terminó mezclando datos de "
        "ATLAS con el runtime de una dependencia externa.",
    ),
    175: (
        "Dentro de esta carpeta, encontramos",
        "El archivo openclaw.json reunía la configuración principal en formato JSON: proveedor, modelo, "
        "canales y extensiones. Sus valores privados se migraron después a un archivo propio de ATLAS, "
        "con permisos restringidos y fuera del repositorio.",
    ),
    176: (
        "Luego, encontramos la carpeta",
        "La carpeta workspace almacenaba las instrucciones y los archivos de conocimiento del agente. "
        "Ese contenido sí pertenecía a ATLAS, por lo que se conservó y se verificó archivo por archivo.",
    ),
    195: ("2.2 Adaptación", "2.2 Adaptación del contexto y migración a .atlas"),
    197: (
        "Una vez visto cómo funciona OpenClaw",
        "La primera adaptación consistió en reescribir los archivos del workspace para convertir un "
        "agente genérico en ATLAS. La segunda, realizada más adelante, fue separar ese conocimiento de "
        "OpenClaw y trasladarlo a una estructura controlada por el propio proyecto.",
    ),
    199: (
        "Además, era importante medir bien",
        "También era importante medir bien el tono. Conservé la idea de escribir instrucciones directas "
        "y humanas, pero desde entonces su organización y su carga ya no dependen del formato interno "
        "de OpenClaw.",
    ),
    205: (
        "Una vez planteado lo que tenía",
        "El archivo principal continuó siendo AGENTS.md, porque explica las relaciones entre el resto "
        "del contexto. Actualmente vive en .atlas/context/knowledge junto con los manuales que carga "
        "WebScreen y atlas-chat.",
    ),
    206: (
        "En MEMORY.md expliqué",
        "MEMORY.md, USER.md y los demás documentos mantienen la memoria y las preferencias duraderas. "
        "Los turnos compartidos se guardan por separado en .atlas/context/conversation para no mezclar "
        "conocimiento estable con el historial que cambia.",
    ),
    207: (
        "En SOUL.md e IDENTITY.md",
        "SOUL.md e IDENTITY.md definen quién es ATLAS y cómo se expresa: sencillo, cercano y adaptado a "
        "la situación, sin perder de vista que forma parte de mi TDR.",
    ),
    208: (
        "Y añadí dos archivos nuevos",
        "ENVIRONMENT.md describe la Raspberry Pi, su hardware y sus límites. TDR.md reúne la información "
        "académica que ATLAS necesita para hablar del proyecto con precisión.",
    ),
    209: (
        "TDR.md era probablemente",
        "TDR.md registra que mi tutor es Rafael Roca Campos, a quien llamamos Rafa; que la parte teórica "
        "se entrega el 21 de octubre; y que la presentación se realizará entre el 9 y el 11 de noviembre "
        "ante un tribunal formado por Rafa y otros dos profesores todavía por determinar. Prepararé una "
        "presentación en Canva y ATLAS me ayudará a explicar el proyecto. También conoce quién es Rafa "
        "para poder saludarlo de forma personalizada cuando esté presente.",
    ),
    221: (
        "Hasta este punto he explicado",
        "Hasta este punto he explicado el origen de ATLAS y el papel que tuvo OpenClaw en el primer "
        "prototipo. El sistema actual ya es un ecosistema propio: combina gpt-realtime-2.1, Native Broker, "
        "contexto y memoria en .atlas, herramientas locales, WebScreen, Companion, la aplicación Android "
        "y el dispositivo físico ATLAS A1.",
    ),
    227: ("3.2 Gateway", "3.2 Native Broker, sesiones y canales de comunicación"),
    228: (
        "OpenClaw utiliza un Gateway",
        "En la arquitectura actual ya no existe un Gateway de OpenClaw. ATLAS Native Broker mantiene un "
        "proceso persistente de codex app-server, comprueba la cuenta, consulta las cuotas y solicita "
        "credenciales efímeras para gpt-realtime-2.1. El OAuth de Codex vive en /home/atlas/.codex/auth.json. "
        "Durante la migración importé y refresqué la sesión existente, por lo que no fue necesario volver "
        "a iniciar sesión.",
    ),
    229: (
        "ATLAS se puede utilizar desde",
        "ATLAS se puede utilizar desde WebScreen, atlas-chat y la aplicación Android. atlas-chat prueba "
        "la misma ruta Realtime desde la terminal, mientras que Companion conecta el S23 Ultra con A1 "
        "mediante Tailscale. El antiguo canal de Telegram se eliminó junto con OpenClaw porque no formaba "
        "parte del flujo que quería para el producto.",
    ),
    230: (
        "Telegram funciona bien",
        "La voz impuso una exigencia distinta a la de un chat asíncrono. El prototipo con OpenClaw podía "
        "producir respuestas de buena calidad, pero normalmente acumulaba entre siete y quince segundos "
        "de espera. La ruta Realtime directa suele situar la primera respuesta hablada alrededor de dos "
        "o tres segundos cuando la sesión ya está preparada.",
    ),
    231: (
        "Las sesiones son otra pieza",
        "Una sesión conserva el contexto inmediato de una conversación. La memoria compartida de los "
        "turnos se guarda en .atlas/context/conversation y el conocimiento estable permanece en "
        ".atlas/context/knowledge. Así, reiniciar una sesión de Realtime no obliga a ATLAS a olvidar su "
        "identidad ni a convertir todo el historial en un bloque infinito.",
    ),
    232: ("3.3 Contexto", "3.3 Contexto, memoria y archivos de .atlas"),
    233: (
        "El workspace sigue siendo",
        "La carpeta .atlas/context/knowledge es el lugar en el que ATLAS entiende su entorno. AGENTS.md "
        "actúa como archivo padre; IDENTITY.md define quién es; SOUL.md determina su carácter; USER.md "
        "contiene preferencias; MEMORY.md resume recuerdos; ENVIRONMENT.md describe A1; TDR.md mantiene "
        "el contexto académico; y TOOLS.md documenta herramientas y convenciones.",
    ),
    235: (
        "La memoria de OpenClaw",
        "La memoria nativa de ATLAS se apoya en Markdown legible y versionable. Esto resulta especialmente "
        "útil en un TDR porque permite ver qué conoce el sistema y corregirlo sin depender de una base de "
        "datos opaca. Un manifest define el orden y la prioridad de carga para WebScreen y atlas-chat.",
    ),
    237: ("3.4 Herramientas, plugins", "3.4 Herramientas, function calling y comandos atlas-*"),
    238: (
        "Un modelo de lenguaje puede explicar",
        "Un modelo puede explicar cómo comprobar el almacenamiento de Linux, pero para decir cuánto queda "
        "realmente necesita observar el sistema. Realtime solicita herramientas nativas como atlas_shell, "
        "atlas_web_search, atlas_routine y las operaciones de Companion; el backend ejecuta la acción y "
        "devuelve un resultado estructurado dentro del mismo turno.",
    ),
    239: (
        "La documentación de OpenClaw separa",
        "El function calling separa razonamiento y ejecución: el modelo decide que necesita una acción, "
        "solicita una función definida por ATLAS y recibe su resultado. Es más verificable que fingir la "
        "respuesta y permite limitar cada herramienta, aunque una operación externa puede seguir añadiendo "
        "latencia al turno.",
    ),
    241: (
        "También separé los archivos",
        "Todo el runtime actual pertenece a .atlas. knowledge contiene instrucciones estables; conversation, "
        "el estado que cambia; roles, los perfiles declarativos; config/secrets.json, únicamente las claves "
        "permitidas con modo 0600; y las demás carpetas, WebScreen, Companion, herramientas y recursos del "
        "dispositivo. OpenClaw queda solo como antecedente histórico y ya no está instalado.",
    ),
    269: (
        "La carpeta .openclaw contiene",
        "La carpeta .atlas reúne el runtime completo: contexto, conversación, roles, Native Broker, "
        "WebScreen, atlas-chat, Companion 0.2.1, rutinas, gestión de pantalla, RAFAS, herramientas y backups. "
        "Las claves de Tavily y ElevenLabs se guardan en .atlas/config/secrets.json con permisos 0600; el "
        "OAuth de Codex se conserva por separado en .codex/auth.json. Ninguna de estas credenciales se "
        "publica en GitHub.",
    ),
    270: (
        "Los servicios se gestionan",
        "Los servicios se gestionan con systemd. WebScreen, Companion, los modos de pantalla, RAFAS y las "
        "conexiones de audio tienen ciclos distintos, por lo que no conviene convertirlos en un único "
        "proceso. Native Broker inicia codex app-server bajo demanda y lo mantiene para evitar pagar el "
        "arranque en cada sesión.",
    ),
    291: (
        "El siguiente cambio fue sustituir",
        "El siguiente cambio fue utilizar la Web Speech API de Chrome como detector local y apoyo para la "
        "entrada del WebScreen. Whisper.cpp se conserva de forma independiente en .atlas/tools como "
        "transcriptor auxiliar y herramienta de diagnóstico; sus binarios y su modelo ya no dependen de "
        "otra aplicación ni participan obligatoriamente en cada turno.",
    ),
    323: (
        "La herramienta atlas_shell ejecuta",
        "La herramienta atlas_shell ejecuta comandos en A1; atlas_web_search busca mediante Tavily; y "
        "atlas_routine administra automatizaciones deterministas. Native Broker prepara la reserva WebRTC "
        "con el OAuth de Codex y entrega a WebScreen o atlas-chat una sesión efímera. No hay un segundo "
        "agente ni una delegación intermedia procesando la conversación.",
    ),
    341: (
        "El rework distingue conexión",
        "El rework distingue la conexión con la Pi, el control del cliente y el transporte Realtime. "
        "WebScreen comprueba el control cada 1,5 segundos, tolera interrupciones breves y conserva una "
        "reserva temporal en el servidor. Un microcorte no debe derribar de inmediato una sesión que aún "
        "puede recuperarse.",
    ),
    343: (
        "El puente comprueba",
        "Native Broker mantiene codex app-server como proceso persistente y reinicia únicamente esa capa "
        "si deja de responder. Los fallos de red, autorización, transcripción y proveedor se registran por "
        "separado. Un HTTP 500 no demuestra por sí solo que el OAuth esté dañado.",
    ),
    372: (
        "RAFAS significa",
        "RAFAS significa Recovery Access For ATLAS Systems. Surgió porque un agente conectado a la nube "
        "no puede repararse si fallan a la vez el proveedor, la red o su interfaz principal. En esas "
        "situaciones ATLAS entra, metafóricamente, en hibernación y necesito una vía física de recuperación.",
    ),
    373: (
        "RAFAS ofrece",
        "RAFAS ofrece una vía local independiente de Chrome, WebScreen, Native Broker y cualquier modelo "
        "en la nube. Al conectar un teclado y pulsar Control + W + O + W, un servicio mínimo enciende la "
        "pantalla si es necesario y cambia a una terminal de recuperación en tty8. También se puede forzar "
        "mediante atlas-screen --rafas.",
    ),
    417: (
        "Las credenciales de proveedores",
        "Las credenciales de proveedores, el OAuth de Codex y las claves de emparejamiento se almacenan "
        "fuera del repositorio público. Las claves permitidas del antiguo prototipo se migraron a "
        ".atlas/config/secrets.json sin mostrarlas, y el archivo quedó limitado al propietario. Cuando se "
        "crea una imagen distribuible, la limpieza se realiza sobre la copia y no sobre A1 en desarrollo.",
    ),
    420: (
        "La fiabilidad se apoya",
        "La fiabilidad se apoya en systemd, atlas-status, atlas-rafas, atlas-chat y los logs. Native Broker "
        "permite comprobar por separado OAuth, cuotas, codex app-server y la reserva de Realtime. Una ruta "
        "de audio ausente, una Pi sin red o un error 500 del proveedor exigen diagnósticos distintos; "
        "volver a autenticar sin comprobar la causa puede no resolver ninguno.",
    ),
    428: (
        "OpenClaw ocupa otra categoría",
        "OpenClaw ocupa otra categoría: es una plataforma de agentes autoalojada, no un asistente doméstico "
        "terminado. Su flexibilidad permitió construir rápidamente el primer prototipo, pero su Gateway y "
        "la delegación acumulaban normalmente entre siete y quince segundos. Por eso lo sustituí por "
        "Native Broker, contexto propio y Realtime directo, que encajan mejor con una conversación física.",
    ),
    455: (
        "Una posibilidad especialmente",
        "Ya he preparado una capa declarativa de roles en .atlas/roles. atlas-full representa el runtime "
        "completo actual. El rol profesores carga únicamente IDENTITY.md y PROFESORES.md, sin shell, red, "
        "escritura ni herramientas, para responder consultas de horarios con una superficie mínima.",
    ),
    456: (
        "El reto será",
        "Los manifiestos y el horario ficticio de diecinueve profesores ya están definidos, pero WebScreen "
        "y atlas-chat todavía funcionan con el equivalente a atlas-full. El selector, la composición de "
        "varios roles y su enforcement en ejecución quedan como trabajo futuro y no se presentan como "
        "funciones terminadas.",
    ),
    483: (
        "OpenClaw 2026.7.1-2",
        "OpenClaw 2026.7.1-2 en el primer prototipo de agente; retirado después de migrar el contexto, las "
        "credenciales necesarias y el flujo Realtime a componentes propios de ATLAS.",
    ),
    491: (
        "Node.js para el puente",
        "Node.js para pruebas de WebScreen, supervisión del navegador y componentes auxiliares; Python para "
        "Native Broker, el backend y la mayor parte de las herramientas del runtime.",
    ),
}


TABLE_REPLACEMENTS: dict[tuple[int, int, int], tuple[str, str]] = {
    (1, 1, 2): (
        "Ejecuta ATLAS OS, OpenClaw",
        "Ejecuta ATLAS OS, Native Broker, WebScreen, Companion y los servicios del dispositivo.",
    ),
    (6, 4, 0): ("Agente OpenClaw", "Prototipo con OpenClaw"),
    (6, 4, 1): ("Memoria, canales", "Integración rápida de memoria, canales y herramientas"),
    (6, 4, 2): (
        "No incluye una interfaz física",
        "Latencia de 7–15 s y una arquitectura externa al producto final",
    ),
    (6, 4, 3): (
        "Contexto, canales y reserva",
        "Antecedente histórico que permitió validar la idea antes del runtime nativo.",
    ),
}


REFINEMENT_REPLACEMENTS: tuple[tuple[str, str], ...] = (
    (
        "Funciona de la siguiente manera. Tras instalarlo, te permite configurar por primera vez "
        "tu agente. Escoges el modelo de IA que quieres usar, los canales de comunicación que "
        "prefieras, diferentes API Keys para poder acceder a recursos en la nube y navegador para "
        "búsquedas, entre muchas otras cosas.",
        "En aquella primera versión el flujo era bastante directo. Tras instalar OpenClaw configuré "
        "el modelo, los canales y las claves de los servicios externos. Esto me permitió tener un "
        "agente funcional muy rápido, aunque también hizo que varias responsabilidades quedasen "
        "agrupadas dentro de una herramienta que no había diseñado yo.",
    ),
    (
        "Una vez completada la configuración, puedes hablar con tu agente mediante uno de los "
        "canales de comunicación, o la página web creada en el localhost.",
        "Después podía hablar con el agente desde los canales configurados o desde una página local. "
        "Para un prototipo era muy cómodo: me permitió comprobar la idea antes de construir WebScreen, "
        "atlas-chat y el resto de la arquitectura propia.",
    ),
    (
        "Así de simple.",
        "Como punto de partida funcionó bien, pero todavía no era el ATLAS independiente que quería "
        "construir.",
    ),
    (
        "Y lo más importante, dentro de dicha carpeta es donde se encuentran varios de los archivos "
        "más importantes para definir la personalidad del agente.",
        "Dentro del workspace de aquella etapa estaban los Markdown que daban identidad y contexto al "
        "agente. Estos archivos sí eran parte de ATLAS y, por eso, fueron una de las piezas que conservé "
        "durante la migración.",
    ),
    (
        "AGENTS.md es leído por el agente siempre al inicio de una sesión, y es el archivo principal "
        "para dotar de contexto al agente. Es, por decirlo así, el archivo padre. Dentro de este, se "
        "explica cual debe ser su funcionamiento, y se referencian todos los otros archivos .md y para "
        "que sirven.",
        "OpenClaw lo cargaba al iniciar una sesión. Más adelante mantuve esa misma relación entre "
        "documentos, pero pasé a resolverla desde el manifest propio de .atlas/context/knowledge.",
    ),
    (
        "Todos estos archivos, son los que permiten al agente (ATLAS, en mi caso), contextualizar su "
        "entorno y situarse.",
        "En resumen, los Markdown eran lo que hacía que un agente genérico entendiese su entorno y se "
        "convirtiese en ATLAS. OpenClaw fue el primer contenedor; el conocimiento era del proyecto.",
    ),
    (
        "No trata al agente como una IA, lo trata como un humano más, y esto lo podemos comprobar con "
        "las primeras dos líneas del archivo principal:",
        "Ese estilo no hablaba al agente como a una máquina, sino como a un colaborador. Se ve muy bien "
        "en las dos primeras líneas del archivo original:",
    ),
    (
        "La voz impuso una exigencia distinta a la de un chat asíncrono. El prototipo con OpenClaw podía "
        "producir respuestas de buena calidad, pero normalmente acumulaba entre siete y quince segundos "
        "de espera. La ruta Realtime directa suele situar la primera respuesta hablada alrededor de dos "
        "o tres segundos cuando la sesión ya está preparada.",
        "La voz impuso una exigencia distinta a la de un chat asíncrono. El prototipo con OpenClaw daba "
        "respuestas de buena calidad, pero en las pruebas de voz acumulaba normalmente entre siete y "
        "quince segundos. En pruebas comparables, la ruta Realtime directa comenzó a hablar "
        "aproximadamente entre dos y tres segundos cuando la sesión ya estaba preparada. Son medidas del "
        "prototipo, no una promesa fija: la red, el audio y las herramientas pueden cambiar el resultado.",
    ),
    (
        "Las credenciales de proveedores, el OAuth de Codex y las claves de emparejamiento se almacenan "
        "fuera del repositorio público. Las claves permitidas del antiguo prototipo se migraron a "
        ".atlas/config/secrets.json sin mostrarlas, y el archivo quedó limitado al propietario. Cuando "
        "se crea una imagen distribuible, la limpieza se realiza sobre la copia y no sobre A1 en desarrollo.",
        "Las claves de Tavily y ElevenLabs se guardan en /home/atlas/.atlas/config/secrets.json, con "
        "permisos 0600 y fuera de Git. El OAuth de Codex permanece separado en "
        "/home/atlas/.codex/auth.json, y las credenciales de emparejamiento se conservan en el estado "
        "privado de Companion. Ninguno de estos datos se imprime en los logs ni se incluye en una imagen "
        "distribuible.",
    ),
    (
        "Crear un índice de contexto para consultar AGENTS.md y sus anexos sin cargarlos completos en "
        "cada sesión.",
        "Evaluar una recuperación selectiva sobre el manifest de contexto existente, para reducir el "
        "tiempo de preparación sin perder instrucciones ni relaciones importantes entre los documentos.",
    ),
    (
        "Ya he preparado una capa declarativa de roles en .atlas/roles. atlas-full representa el runtime "
        "completo actual. El rol profesores carga únicamente IDENTITY.md y PROFESORES.md, sin shell, red, "
        "escritura ni herramientas, para responder consultas de horarios con una superficie mínima.",
        "Implementar un selector de roles que valide cada manifest y aplique sus capacidades antes de "
        "iniciar la sesión, sin conceder permisos por el simple hecho de cargar un Markdown.",
    ),
    (
        "Los manifiestos y el horario ficticio de diecinueve profesores ya están definidos, pero "
        "WebScreen y atlas-chat todavía funcionan con el equivalente a atlas-full. El selector, la "
        "composición de varios roles y su enforcement en ejecución quedan como trabajo futuro y no se "
        "presentan como funciones terminadas.",
        "Probar la composición explícita entre atlas-full y profesores, manteniendo profesores como "
        "perfil cerrado cuando se use solo y dejando claro en la interfaz qué contexto y qué herramientas "
        "están activos.",
    ),
    (
        "Python para el backend de WebScreen y utilidades del sistema.",
        "Python para Native Broker, el backend de WebScreen, el motor de rutinas y la mayor parte de las "
        "herramientas del runtime.",
    ),
    (
        "Node.js para pruebas de WebScreen, supervisión del navegador y componentes auxiliares; Python "
        "para Native Broker, el backend y la mayor parte de las herramientas del runtime.",
        "Node.js para pruebas de WebScreen, supervisión del navegador y componentes auxiliares de "
        "desarrollo; ya no existe un puente persistente con OpenClaw en el producto actual.",
    ),
    (
        "La sesión puede mantener contexto durante la conversación y admite varias voces nativas. La "
        "interfaz incluye un selector con alloy, ash, ballad, cedar, coral, echo, marin, sage, shimmer y "
        "verse. Una limitación relevante es que la voz no puede cambiarse una vez que la sesión ya ha "
        "emitido audio; para aplicar el cambio se crea una sesión nueva.",
        "Realtime admite voces nativas como alloy, ash, ballad, beacon, cedar, coral, echo, marin, sage, "
        "shimmer y verse, entre otras. Una limitación relevante es que la voz no puede cambiarse una vez "
        "que la sesión ya ha emitido audio; para aplicar el cambio se crea una sesión nueva.",
    ),
    (
        "El hecho de que la lista sea larga no significa que ATLAS sea una simple suma de programas. "
        "El trabajo del proyecto está en integrarlos, decidir qué responsabilidad tiene cada uno y "
        "construir una experiencia coherente encima. Muchas herramientas podrían sustituirse; la "
        "arquitectura y las decisiones que las conectan son lo que define el sistema.",
        "La lista no convierte ATLAS en una simple suma de programas. El proyecto consiste en integrarlos, "
        "asignar responsabilidades y construir una experiencia coherente. Muchas herramientas podrían "
        "sustituirse; lo que define el sistema es la arquitectura que las conecta.",
    ),
)


TDR_COMBINED = (
    "TDR.md registra que mi tutor es Rafael Roca Campos, a quien llamamos Rafa; que la parte teórica se "
    "entrega el 21 de octubre; y que la presentación se realizará entre el 9 y el 11 de noviembre ante "
    "un tribunal formado por Rafa y otros dos profesores todavía por determinar. Prepararé una "
    "presentación en Canva y ATLAS me ayudará a explicar el proyecto. También conoce quién es Rafa para "
    "poder saludarlo de forma personalizada cuando esté presente."
)
TDR_DATES = (
    "TDR.md guarda las dos fechas que ATLAS debe tener presentes: la parte teórica se entrega el 21 de "
    "octubre y la presentación se hará entre el 9 y el 11 de noviembre."
)
TDR_PRESENTATION = (
    "Mi tutor es Rafael Roca Campos, a quien llamamos Rafa. La defensa será ante un tribunal de tres "
    "profesores: Rafa y otros dos que todavía no conozco. Prepararé una presentación en Canva y quiero "
    "que el propio ATLAS me ayude a explicar el proyecto; por eso también conoce quién es Rafa y puede "
    "dirigirse a él de forma natural cuando esté presente."
)

BROKER_OLD = (
    "En la arquitectura actual ya no existe un Gateway de OpenClaw. ATLAS Native Broker mantiene un "
    "proceso persistente de codex app-server, comprueba la cuenta, consulta las cuotas y solicita "
    "credenciales efímeras para gpt-realtime-2.1. El OAuth de Codex vive en "
    "/home/atlas/.codex/auth.json. Durante la migración importé y refresqué la sesión existente, por lo "
    "que no fue necesario volver a iniciar sesión."
)
BROKER_REFINED = (
    "En la arquitectura actual ya no existe un Gateway de OpenClaw. ATLAS Native Broker mantiene un "
    "proceso codex app-server durante la vida del servicio, comprueba el estado de la cuenta, normaliza "
    "las cuotas y crea secretos efímeros para las sesiones de gpt-realtime-2.1. El OAuth persistente "
    "sigue siendo propiedad de Codex y se conserva en /home/atlas/.codex/auth.json; el broker no lo "
    "copia a .atlas, no lo devuelve por HTTP y no lo escribe en los logs."
)
BROKER_DIAGNOSTICS = (
    "Para diagnosticar esta capa creé atlas-broker health, atlas-broker usage y atlas-broker session. "
    "Cada comando comprueba una parte distinta sin imprimir tokens ni identificadores: el proceso y el "
    "login, las ventanas de uso o la capacidad de crear una reserva Realtime."
)

RUNTIME_PARAGRAPH = (
    "Todo el runtime actual pertenece a .atlas. knowledge contiene instrucciones estables; conversation, "
    "el estado que cambia; roles, los perfiles declarativos; config/secrets.json, únicamente las claves "
    "permitidas con modo 0600; y las demás carpetas, WebScreen, Companion, herramientas y recursos del "
    "dispositivo. OpenClaw queda solo como antecedente histórico y ya no está instalado."
)
ROLE_CURRENT = (
    "También he separado el contexto por roles. atlas-full describe el ATLAS completo y equivale al "
    "comportamiento actual. profesores es una demostración de mínima autoridad: carga únicamente "
    "IDENTITY.md y PROFESORES.md, recibe la fecha y la hora actuales y no puede usar shell, red, escritura, "
    "memoria mutable ni herramientas del dispositivo."
)
ROLE_FUTURE = (
    "Los manifiestos ya están versionados, pero todavía no existe un selector dinámico ni un sistema de "
    "composición activo. El horario de diecinueve profesores es ficticio y sirve como demostración; por "
    "tanto, esta parte prueba cómo limitar contexto y capacidades, no representa todavía un despliegue "
    "real del instituto."
)

NUMBERED_REFINEMENTS = (
    "Ampliar la validación física en el S23 Ultra con BLE, biometría, llamadas y SMS únicamente cuando "
    "se disponga de condiciones seguras y confirmación explícita.",
    "Repetir el bucle de control visual con tareas reproducibles en varias aplicaciones, medir cada paso "
    "y verificar que la parada interrumpe cualquier secuencia sin conservar capturas privadas.",
    "Definir ACL y caducidad de dispositivos en Tailscale, rotación de claves de emparejamiento y una ruta "
    "de recuperación que no vuelva a exponer el Companion públicamente.",
)

REFINED_TABLE_ROW = (
    "Primer prototipo con OpenClaw",
    "Permitió unir rápidamente modelo, memoria, canales y herramientas.",
    "En las pruebas de voz acumulaba normalmente entre 7 y 15 s y dependía de un Gateway externo a ATLAS.",
    "Antecedente histórico que validó la idea; no forma parte del runtime actual.",
)


def replace_paragraph(paragraph, text: str) -> None:
    """Replace visible paragraph text while preserving its first run formatting."""
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            paragraph._p.remove(run._r)
        for hyperlink in list(paragraph._p.xpath("./w:hyperlink")):
            paragraph._p.remove(hyperlink)
    else:
        paragraph.add_run(text)


def _normalized(text: str) -> str:
    return " ".join(text.split())


def _paragraphs_with_text(document, text: str) -> list[Paragraph]:
    expected = _normalized(text)
    return [paragraph for paragraph in document.paragraphs if _normalized(paragraph.text) == expected]


def _find_unique_paragraph(document, text: str) -> Paragraph:
    matches = _paragraphs_with_text(document, text)
    if len(matches) != 1:
        raise SystemExit(
            f"Esperaba un párrafo con el texto {text!r}, pero encontré {len(matches)} coincidencias."
        )
    return matches[0]


def _replace_once(document, old: str, new: str) -> bool:
    """Replace an exact body paragraph once; accept an already-refined document."""
    if _paragraphs_with_text(document, new):
        return False
    paragraph = _find_unique_paragraph(document, old)
    replace_paragraph(paragraph, new)
    return True


def _replace_agents_paragraph(document) -> bool:
    """Refine P187 while retaining its AGENTS.md hyperlink and its run formatting."""
    old = (
        "Y el archivo principal y más importante: “AGENTS.md”. Este define todo el funcionamiento del "
        "agente."
    )
    new = (
        "El archivo principal de aquel conjunto era AGENTS.md, porque enlazaba y explicaba el resto del "
        "contexto."
    )
    if _paragraphs_with_text(document, new):
        return False
    paragraph = _find_unique_paragraph(document, old)
    hyperlinks = paragraph._p.xpath("./w:hyperlink")
    matching = [
        hyperlink
        for hyperlink in hyperlinks
        if "".join(hyperlink.xpath(".//w:t/text()")).strip() == "AGENTS.md"
    ]
    if len(matching) != 1:
        raise SystemExit("No pude conservar de forma inequívoca el enlace de AGENTS.md en P187.")

    direct_runs = paragraph._p.xpath("./w:r")
    run_properties = None
    if direct_runs:
        candidate = direct_runs[0].find(qn("w:rPr"))
        if candidate is not None:
            run_properties = deepcopy(candidate)
    hyperlink = deepcopy(matching[0])

    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    before = paragraph.add_run("El archivo principal de aquel conjunto era ")
    if run_properties is not None:
        before._r.insert(0, deepcopy(run_properties))
    paragraph._p.append(hyperlink)
    after = paragraph.add_run(", porque enlazaba y explicaba el resto del contexto.")
    if run_properties is not None:
        after._r.insert(0, deepcopy(run_properties))
    return True


def _insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_element = OxmlElement("w:p")
    if paragraph._p.pPr is not None:
        new_element.append(deepcopy(paragraph._p.pPr))
    paragraph._p.addnext(new_element)
    inserted = Paragraph(new_element, paragraph._parent)
    run = inserted.add_run(text)
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(paragraph.runs[0]._r.rPr))
    return inserted


def _ensure_sequence_after(document, anchor_text: str, texts: tuple[str, ...]) -> int:
    """Insert missing paragraphs in order directly after an exact anchor."""
    cursor = _find_unique_paragraph(document, anchor_text)
    inserted_count = 0
    for text in texts:
        existing = _paragraphs_with_text(document, text)
        if existing:
            if len(existing) != 1:
                raise SystemExit(f"El párrafo refinado aparece repetido: {text!r}")
            cursor = existing[0]
            continue
        cursor = _insert_paragraph_after(cursor, text)
        inserted_count += 1
    return inserted_count


def _ensure_numbering(document, text: str, *, num_id: int = 2, level: int = 0) -> bool:
    paragraph = _find_unique_paragraph(document, text)
    properties = paragraph._p.get_or_add_pPr()
    current = properties.numPr
    if current is not None:
        current_num = current.numId.val if current.numId is not None else None
        current_level = current.ilvl.val if current.ilvl is not None else None
        if current_num == num_id and current_level == level:
            return False
        properties.remove(current)

    numbering = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numbering.extend((ilvl, num))
    properties.append(numbering)
    return True


def _remove_red_run_colors(document) -> int:
    """Remove explicit pure-red run colors from every XML part in the package."""
    removed = 0
    for part in document.part.package.parts:
        element = getattr(part, "_element", None)
        if element is None:
            continue
        for color in list(element.xpath(".//w:rPr/w:color")):
            value = (color.get(qn("w:val")) or "").upper()
            if value == "FF0000":
                color.getparent().remove(color)
                removed += 1
    return removed


def _replace_table_row(document) -> int:
    row = document.tables[6].rows[4]
    if tuple(_normalized(cell.text) for cell in row.cells) == tuple(
        _normalized(value) for value in REFINED_TABLE_ROW
    ):
        return 0
    for cell, value in zip(row.cells, REFINED_TABLE_ROW, strict=True):
        replace_paragraph(cell.paragraphs[0], value)
        for extra in cell.paragraphs[1:]:
            replace_paragraph(extra, "")
    return len(REFINED_TABLE_ROW)


def _save_atomically(document, destination: Path) -> None:
    temporary = destination.with_name(f".{destination.name}.refining.tmp")
    try:
        document.save(temporary)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)


def refine() -> None:
    """Apply the post-migration editorial pass safely to the current DOCX."""
    if not DOCUMENT.is_file():
        raise SystemExit(f"No existe el documento: {DOCUMENT}")
    document = Document(DOCUMENT)
    changes = 0

    changes += int(_replace_agents_paragraph(document))
    for old, new in REFINEMENT_REPLACEMENTS:
        changes += int(_replace_once(document, old, new))

    if not _paragraphs_with_text(document, TDR_DATES):
        changes += int(_replace_once(document, TDR_COMBINED, TDR_DATES))
    changes += _ensure_sequence_after(document, TDR_DATES, (TDR_PRESENTATION,))

    if not _paragraphs_with_text(document, BROKER_REFINED):
        changes += int(_replace_once(document, BROKER_OLD, BROKER_REFINED))
    changes += _ensure_sequence_after(document, BROKER_REFINED, (BROKER_DIAGNOSTICS,))
    changes += _ensure_sequence_after(document, RUNTIME_PARAGRAPH, (ROLE_CURRENT, ROLE_FUTURE))

    for text in NUMBERED_REFINEMENTS:
        changes += int(_ensure_numbering(document, text))

    changes += _replace_table_row(document)
    changes += _remove_red_run_colors(document)

    if changes:
        _save_atomically(document, DOCUMENT)
        print(f"Refinamiento aplicado: {changes} cambios estructurales o de formato.")
    else:
        print("El refinamiento ya estaba aplicado; el documento no se reescribió.")


def migrate() -> None:
    if not DOCUMENT.is_file():
        raise SystemExit(f"No existe el documento: {DOCUMENT}")
    document = Document(DOCUMENT)
    for index, (expected, replacement) in REPLACEMENTS.items():
        paragraph = document.paragraphs[index]
        if expected.casefold() not in paragraph.text.casefold():
            raise SystemExit(
                f"El párrafo {index} cambió: esperaba {expected!r}, encontré {paragraph.text!r}"
            )
        replace_paragraph(paragraph, replacement)

    for (table_index, row_index, cell_index), (expected, replacement) in TABLE_REPLACEMENTS.items():
        cell = document.tables[table_index].rows[row_index].cells[cell_index]
        if expected.casefold() not in cell.text.casefold():
            raise SystemExit(
                f"La celda T{table_index}R{row_index}C{cell_index} cambió: "
                f"esperaba {expected!r}, encontré {cell.text!r}"
            )
        replace_paragraph(cell.paragraphs[0], replacement)
        for extra in cell.paragraphs[1:]:
            replace_paragraph(extra, "")

    if not BACKUP.exists():
        shutil.copy2(DOCUMENT, BACKUP)
    document.save(DOCUMENT)
    print(f"Actualizados {len(REPLACEMENTS)} párrafos y {len(TABLE_REPLACEMENTS)} celdas.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--refine",
        action="store_true",
        help="aplica el refinamiento editorial idempotente sobre el DOCX ya migrado",
    )
    args = parser.parse_args()
    if args.refine:
        refine()
    else:
        migrate()


if __name__ == "__main__":
    main()
